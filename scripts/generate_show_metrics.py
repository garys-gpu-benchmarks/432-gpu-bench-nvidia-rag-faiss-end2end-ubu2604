#!/usr/bin/env python3
# File: scripts/generate_show_metrics.py
# Version: 1.1.0
# Maintainer: AI Agent GPU Benchmark Repo Generator
# Date: 2026-10-11
# Description: Generate the two built-in tables of
#   templates/suite-tools/scripts/show_metrics.sh from BenchmarkSpecDefinitions.xlsx:
#
#   Metric_Short_Labels   one row per workload metric (Workload Number, Metric No,
#                         Metric Key, Short Label <= 7 characters, ...)
#                         -> block "# BEGIN/END GENERATED SHORT LABELS"
#   Workload_Short_Names  one row per workload (Workload Number, Workload Name,
#                         Short Name <= 12 characters, ...)
#                         -> block "# BEGIN/END GENERATED WORKLOAD NAMES"
#
#   Both sheets are checked against Workload_Definitions (metrics parsed as
#   update_runtime_ledger.py and show_metrics.sh do) before anything is written.
#
#   A metric key whose label is the same in every workload is written once
#   ("key label"); a workload whose label differs from the key's usual label
#   gets its own line ("WORKLOAD:key label"), which show_metrics.sh prefers.
#
# Execution:
#   python3 scripts/generate_show_metrics.py              # rewrite the template
#   python3 scripts/generate_show_metrics.py --check      # exit 1 if out of date or a sheet is invalid
#   python3 scripts/generate_show_metrics.py --add-missing
#       add rows for metrics and workloads that have none (metric label from the
#       same key elsewhere, else an automatic abbreviation; workload name from
#       the same full name elsewhere, else its first 12 characters), marked for
#       review, then rewrite
# Requirements: Python 3.10+, openpyxl
# License: Apache-2.0

from __future__ import annotations

import argparse
import collections
import re
import sys
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_NAME = "BenchmarkSpecDefinitions.xlsx"
DEFINITIONS_SHEET = "Workload_Definitions"
LABEL_SHEET = "Metric_Short_Labels"
NAME_SHEET = "Workload_Short_Names"
TEMPLATE_REL = Path("templates") / "suite-tools" / "scripts" / "show_metrics.sh"
BEGIN_MARKER = "# BEGIN GENERATED SHORT LABELS"
END_MARKER = "# END GENERATED SHORT LABELS"
NAME_BEGIN_MARKER = "# BEGIN GENERATED WORKLOAD NAMES"
NAME_END_MARKER = "# END GENERATED WORKLOAD NAMES"

MAX_LABEL = 7  # show_metrics.sh prints every metric in a 7-character column
MAX_NAME = 12  # show_metrics.sh Name column; with the Exit column and 7 metrics a line is then exactly 93 characters
METRIC_SLOTS = 7  # update_runtime_ledger.METRIC_SLOT_COUNT
# Labels are read by awk as the last whitespace-separated word of a line and
# may also be given in a CSV, so no spaces, tabs, commas, colons or quotes.
LABEL_RE = re.compile(r"^[A-Za-z0-9_%/.+-]+$")
# Workload short names may contain single spaces; no tabs, commas, quotes or colons.
NAME_RE = re.compile(r"^[A-Za-z0-9()+&/._-]+(?: [A-Za-z0-9()+&/._-]+)*$")

LABEL_HEADERS = [
    "Workload Number",
    "Workload Name",
    "GPU Vendor",
    "Metric No",
    "Metric Key",
    "Metric Description",
    "Short Label",
    "Label Length",
    "Label Note",
]
NAME_HEADERS = [
    "Workload Number",
    "Workload Name",
    "GPU Vendor",
    "Short Name",
    "Name Length",
    "Name Note",
]
AUTO_NOTE = "auto: review"


class SheetError(Exception):
    """The workbook or the label sheet cannot be used."""


@dataclass(frozen=True)
class Metric:
    workload: int
    workload_name: str
    vendor: str
    number: int
    key: str
    description: str


@dataclass
class LabelRow:
    row: int
    workload: int
    number: int
    key: str
    label: str


@dataclass(frozen=True)
class Workload:
    number: int
    name: str
    vendor: str


@dataclass
class NameRow:
    row: int
    workload: int
    full_name: str
    short: str


# ------------------------------------------------------------------------------
# Metrics parsing (must match update_runtime_ledger.metric_items and the
# metric_key() function in show_metrics.sh)
# ------------------------------------------------------------------------------
def split_metrics(text: str) -> list[tuple[int, str]]:
    """'#1: A (a); #2: B (b);' -> [(1, 'A (a)'), (2, 'B (b)')]."""
    items: list[tuple[int, str]] = []
    for part in re.split(r";\s*|,\s*(?=#\d)|\n", text or ""):
        match = re.match(r"^#(\d+)\s*:\s*(.*)$", part.strip())
        if match:
            items.append((int(match.group(1)), match.group(2).strip()))
    return items


def metric_key(definition: str) -> str:
    """Text in the last (...) at the end; else the first (snake_case); else the text."""
    match = re.search(r"\(([^()]*)\)\s*$", definition)
    if match:
        return match.group(1).strip()
    match = re.search(r"\(([A-Za-z0-9_]+)\)", definition)
    if match:
        return match.group(1).strip()
    return definition.strip()


def metric_description(definition: str, key: str) -> str:
    text = definition.replace(f"({key})", " ")
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s+,", ",", text)
    return text.strip(" ,;")


def auto_label(key: str) -> str:
    """Fallback abbreviation, same rule as auto_label() in show_metrics.sh."""
    return re.sub(r"[^A-Za-z0-9%/]", "", key)[:MAX_LABEL] or "M"


# ------------------------------------------------------------------------------
# Workbook access
# ------------------------------------------------------------------------------
def _norm(value: object) -> str:
    return "" if value is None else str(value).strip()


def _as_int(value: object) -> int | None:
    text = _norm(value)
    if re.fullmatch(r"\d+(\.0+)?", text):
        return int(float(text))
    return None


def _header_index(header: list[object], name: str) -> int:
    wanted = name.lower()
    for index, cell in enumerate(header):
        if _norm(cell).lower() == wanted:
            return index
    raise SheetError(f"column '{name}' not found")


def load_workbook_file(path: Path, read_only: bool = True):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on host
        raise SheetError("openpyxl is required (pip install openpyxl)") from exc
    if not path.is_file():
        raise SheetError(f"workbook not found: {path}")
    return load_workbook(path, read_only=read_only, data_only=False)


def workbook_metrics(workbook) -> list[Metric]:
    """Every metric of every workload in Workload_Definitions, in sheet order."""
    if DEFINITIONS_SHEET not in workbook.sheetnames:
        raise SheetError(f"sheet '{DEFINITIONS_SHEET}' not found")
    rows = list(workbook[DEFINITIONS_SHEET].iter_rows(values_only=True))
    if not rows:
        raise SheetError(f"sheet '{DEFINITIONS_SHEET}' is empty")
    header = list(rows[0])
    c_num = _header_index(header, "Workload Number")
    c_name = _header_index(header, "Workload Name")
    c_vendor = _header_index(header, "GPU Vendor")
    c_metrics = _header_index(header, "Metrics")
    metrics: list[Metric] = []
    for row in rows[1:]:
        workload = _as_int(row[c_num] if c_num < len(row) else None)
        if workload is None:  # blank rows and repeated header rows
            continue
        for number, definition in split_metrics(_norm(row[c_metrics])):
            key = metric_key(definition)
            metrics.append(Metric(
                workload=workload,
                workload_name=_norm(row[c_name]),
                vendor=_norm(row[c_vendor]),
                number=number,
                key=key,
                description=metric_description(definition, key),
            ))
    return metrics


def label_rows(workbook) -> list[LabelRow]:
    if LABEL_SHEET not in workbook.sheetnames:
        raise SheetError(f"sheet '{LABEL_SHEET}' not found in {WORKBOOK_NAME}")
    rows = list(workbook[LABEL_SHEET].iter_rows(values_only=True))
    if not rows:
        raise SheetError(f"sheet '{LABEL_SHEET}' is empty")
    header = list(rows[0])
    c_num = _header_index(header, "Workload Number")
    c_no = _header_index(header, "Metric No")
    c_key = _header_index(header, "Metric Key")
    c_label = _header_index(header, "Short Label")
    out: list[LabelRow] = []
    for index, row in enumerate(rows[1:], start=2):
        def cell(col: int) -> object:
            return row[col] if col < len(row) else None
        workload = _as_int(cell(c_num))
        if workload is None and not any(_norm(v) for v in row):
            continue
        out.append(LabelRow(
            row=index,
            workload=workload if workload is not None else -1,
            number=_as_int(cell(c_no)) or -1,
            key=_norm(cell(c_key)),
            label=_norm(cell(c_label)),
        ))
    return out


def workbook_workloads(workbook) -> list[Workload]:
    """Every workload in Workload_Definitions (including any with no metrics), in sheet order."""
    rows = list(workbook[DEFINITIONS_SHEET].iter_rows(values_only=True))
    header = list(rows[0])
    c_num = _header_index(header, "Workload Number")
    c_name = _header_index(header, "Workload Name")
    c_vendor = _header_index(header, "GPU Vendor")
    seen: dict[int, Workload] = {}
    for row in rows[1:]:
        number = _as_int(row[c_num] if c_num < len(row) else None)
        if number is not None and number not in seen:
            seen[number] = Workload(number, _norm(row[c_name]), _norm(row[c_vendor]))
    return list(seen.values())


def name_rows(workbook) -> list[NameRow]:
    if NAME_SHEET not in workbook.sheetnames:
        raise SheetError(f"sheet '{NAME_SHEET}' not found in {WORKBOOK_NAME}; run with --add-missing")
    rows = list(workbook[NAME_SHEET].iter_rows(values_only=True))
    if not rows:
        raise SheetError(f"sheet '{NAME_SHEET}' is empty")
    header = list(rows[0])
    c_num = _header_index(header, "Workload Number")
    c_name = _header_index(header, "Workload Name")
    c_short = _header_index(header, "Short Name")
    out: list[NameRow] = []
    for index, row in enumerate(rows[1:], start=2):
        def cell(col: int) -> object:
            return row[col] if col < len(row) else None
        workload = _as_int(cell(c_num))
        if workload is None and not any(_norm(v) for v in row):
            continue
        out.append(NameRow(row=index, workload=workload if workload is not None else -1,
                           full_name=_norm(cell(c_name)), short=_norm(cell(c_short))))
    return out


def auto_name(full_name: str) -> str:
    """Fallback short name: the first 12 characters, the same as show_metrics.sh without a table entry."""
    text = re.sub(r"[^A-Za-z0-9()+&/._ -]", "", full_name)
    return re.sub(r" +", " ", text)[:MAX_NAME].strip() or "Workload"


def validate_names(workloads: list[Workload], names: list[NameRow]) -> list[str]:
    errors: list[str] = []
    expected = {w.number: w for w in workloads}
    seen: dict[int, int] = {}
    in_batch: dict[tuple[int, str], int] = {}
    by_full: dict[str, tuple[str, int]] = {}

    for item in names:
        where = f"{NAME_SHEET} row {item.row}"
        if item.workload < 0:
            errors.append(f"{where}: Workload Number must be a whole number")
            continue
        if item.workload in seen:
            errors.append(f"{where}: workload {item.workload} is also on row {seen[item.workload]}")
            continue
        seen[item.workload] = item.row
        workload = expected.get(item.workload)
        if workload is None:
            errors.append(f"{where}: workload {item.workload} is not in {DEFINITIONS_SHEET}")
            continue
        if item.full_name != workload.name:
            errors.append(f"{where}: Workload Name '{item.full_name}' does not match {DEFINITIONS_SHEET} "
                          f"('{workload.name}'); update the row")
        if not item.short:
            errors.append(f"{where}: Short Name is empty")
            continue
        if len(item.short) > MAX_NAME:
            errors.append(f"{where}: Short Name '{item.short}' is longer than {MAX_NAME} characters")
        if not NAME_RE.match(item.short):
            errors.append(f"{where}: Short Name '{item.short}' may use letters, digits, single spaces "
                          "and ( ) + & / . _ - only")
        batch = (item.workload // 100, item.short.lower())
        if batch in in_batch:
            errors.append(f"{where}: Short Name '{item.short}' is also used by workload {in_batch[batch]} "
                          f"in the {item.workload // 100}xx batch")
        in_batch[batch] = item.workload
        previous = by_full.get(workload.name)
        if previous and previous[0] != item.short:
            errors.append(f"{where}: '{workload.name}' is '{previous[0]}' for workload {previous[1]} but "
                          f"'{item.short}' here; use one short name per full name")
        by_full.setdefault(workload.name, (item.short, item.workload))

    for number, workload in expected.items():
        if number not in seen:
            errors.append(f"{NAME_SHEET}: no row for workload {number} ({workload.name}); run with --add-missing")
    return errors


def build_name_table(names: list[NameRow]) -> list[tuple[str, str]]:
    return [(str(r.workload), r.short) for r in sorted(names, key=lambda r: r.workload)]


# ------------------------------------------------------------------------------
# Validation and table building
# ------------------------------------------------------------------------------
def validate(metrics: list[Metric], labels: list[LabelRow]) -> list[str]:
    errors: list[str] = []
    expected = {(m.workload, m.number): m for m in metrics}
    seen: dict[tuple[int, int], int] = {}
    per_workload: dict[int, dict[str, int]] = collections.defaultdict(dict)

    for metric in metrics:
        if metric.number > METRIC_SLOTS:
            errors.append(f"workload {metric.workload} metric #{metric.number}: the ledger keeps only "
                          f"{METRIC_SLOTS} metrics")

    for item in labels:
        where = f"{LABEL_SHEET} row {item.row}"
        if item.workload < 0 or item.number < 0:
            errors.append(f"{where}: Workload Number and Metric No must be whole numbers")
            continue
        ident = (item.workload, item.number)
        if ident in seen:
            errors.append(f"{where}: workload {item.workload} metric #{item.number} is also on row {seen[ident]}")
            continue
        seen[ident] = item.row
        metric = expected.get(ident)
        if metric is None:
            errors.append(f"{where}: workload {item.workload} has no metric #{item.number} in {DEFINITIONS_SHEET}")
            continue
        if item.key != metric.key:
            errors.append(f"{where}: Metric Key '{item.key}' does not match {DEFINITIONS_SHEET} "
                          f"('{metric.key}') for workload {item.workload} metric #{item.number}")
        if not item.label:
            errors.append(f"{where}: Short Label is empty")
            continue
        if len(item.label) > MAX_LABEL:
            errors.append(f"{where}: Short Label '{item.label}' is longer than {MAX_LABEL} characters")
        if not LABEL_RE.match(item.label):
            errors.append(f"{where}: Short Label '{item.label}' may use only letters, digits and _ % / . + -")
        other = per_workload[item.workload].get(item.label)
        if other is not None:
            errors.append(f"{where}: Short Label '{item.label}' is used twice in workload {item.workload} "
                          f"(also metric #{other})")
        per_workload[item.workload][item.label] = item.number

    for ident, metric in expected.items():
        if ident not in seen:
            errors.append(f"{LABEL_SHEET}: no row for workload {metric.workload} metric #{metric.number} "
                          f"({metric.key}); run with --add-missing")
    return errors


def build_table(labels: list[LabelRow]) -> list[tuple[str, str]]:
    """(lookup key, label) lines: one per key, plus WORKLOAD:key where a workload differs."""
    ordered = sorted(labels, key=lambda r: (r.workload, r.number))
    by_key: dict[str, collections.Counter] = collections.defaultdict(collections.Counter)
    first_seen: dict[str, int] = {}
    for position, item in enumerate(ordered):
        by_key[item.key][item.label] += 1
        first_seen.setdefault(item.key, position)

    default: dict[str, str] = {}
    for key, counts in by_key.items():
        top = max(counts.values())
        # Most common label; on a tie, the one used by the lowest workload.
        default[key] = next(r.label for r in ordered if r.key == key and counts[r.label] == top)

    lines = [(key, default[key]) for key in sorted(default, key=first_seen.__getitem__)]
    for item in ordered:
        if item.label != default[item.key]:
            lines.append((f"{item.workload}:{item.key}", item.label))
    return lines


@dataclass(frozen=True)
class Block:
    begin: str
    end: str
    sheet: str
    noun: str
    min_width: int


LABEL_BLOCK = Block(BEGIN_MARKER, END_MARKER, LABEL_SHEET, "labels", 57)
NAME_BLOCK = Block(NAME_BEGIN_MARKER, NAME_END_MARKER, NAME_SHEET, "workload short names", 6)


def render_block(lines: list[tuple[str, str]], indent: str, block: Block = LABEL_BLOCK) -> str:
    width = max((len(k) for k, _ in lines), default=0) + 2
    width = max(width, block.min_width)
    body = "".join(f"{key.ljust(width)}{value}\n" for key, value in lines)
    return (
        f"{indent}{block.begin}\n"
        f"{indent}# {len(lines)} {block.noun} from the {block.sheet} sheet of {WORKBOOK_NAME}\n"
        f"{indent}# (scripts/generate_show_metrics.py in the repository generator)\n"
        f"{indent}cat <<'EOF'\n"
        f"{body}"
        "EOF\n"
        f"{indent}{block.end}\n"
    )


def apply_block(text: str, lines: list[tuple[str, str]], block: Block = LABEL_BLOCK) -> str:
    pattern = re.compile(
        r"^(?P<indent>[ \t]*)" + re.escape(block.begin) + r"\n.*?^[ \t]*" + re.escape(block.end) + r"\n",
        flags=re.MULTILINE | re.DOTALL,
    )
    match = pattern.search(text)
    if match is None:
        raise SheetError(f"{TEMPLATE_REL}: markers '{block.begin}' / '{block.end}' not found")
    return text[: match.start()] + render_block(lines, match.group("indent"), block) + text[match.end():]


# ------------------------------------------------------------------------------
# --add-missing
# ------------------------------------------------------------------------------
def _new_sheet(workbook, title: str, headers: list[str], column: str, limit: int, what: str):
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    sheet = workbook.create_sheet(title)  # last, after Workload_Definitions
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    validation = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=str(limit),
                                showErrorMessage=True, errorTitle=what,
                                error=f"{what}s are at most {limit} characters.")
    sheet.add_data_validation(validation)
    validation.add(f"{column}2:{column}5000")
    return sheet


def add_missing_rows(workbook_path: Path) -> tuple[int, int]:
    """Append rows for metrics and workloads with none (creating the sheets if needed).

    Returns (metric rows added, workload rows added).
    """
    workbook = load_workbook_file(workbook_path, read_only=False)
    metrics = workbook_metrics(workbook)
    workloads = workbook_workloads(workbook)

    # Metric_Short_Labels
    if LABEL_SHEET in workbook.sheetnames:
        sheet = workbook[LABEL_SHEET]
        existing = label_rows(workbook)
    else:
        sheet = _new_sheet(workbook, LABEL_SHEET, LABEL_HEADERS, "G", MAX_LABEL, "Short Label")
        existing = []
    have = {(r.workload, r.number) for r in existing}
    known = collections.defaultdict(collections.Counter)
    for r in existing:
        if r.label:
            known[r.key][r.label] += 1
    labels_added = 0
    for metric in metrics:
        if (metric.workload, metric.number) in have:
            continue
        if known[metric.key]:
            label, note = known[metric.key].most_common(1)[0][0], "same label as this key elsewhere"
        else:
            label, note = auto_label(metric.key), AUTO_NOTE
        row = sheet.max_row + 1
        sheet.append([metric.workload, metric.workload_name, metric.vendor, metric.number, metric.key,
                      metric.description, label, f"=LEN(G{row})", note])
        known[metric.key][label] += 1
        labels_added += 1

    # Workload_Short_Names
    if NAME_SHEET in workbook.sheetnames:
        sheet = workbook[NAME_SHEET]
        existing_names = name_rows(workbook)
    else:
        sheet = _new_sheet(workbook, NAME_SHEET, NAME_HEADERS, "D", MAX_NAME, "Short Name")
        existing_names = []
    have_names = {r.workload for r in existing_names}
    by_full = {r.full_name: r.short for r in existing_names if r.short}
    names_added = 0
    for workload in workloads:
        if workload.number in have_names:
            continue
        if workload.name in by_full:
            short, note = by_full[workload.name], "same short name as this full name elsewhere"
        else:
            short, note = auto_name(workload.name), AUTO_NOTE
        row = sheet.max_row + 1
        sheet.append([workload.number, workload.name, workload.vendor, short, f"=LEN(D{row})", note])
        by_full.setdefault(workload.name, short)
        names_added += 1

    if labels_added or names_added:
        workbook.save(workbook_path)
    return labels_added, names_added


# ------------------------------------------------------------------------------
# main
# ------------------------------------------------------------------------------
def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Generate show_metrics.sh tables from Metric_Short_Labels and Workload_Short_Names")
    parser.add_argument("--template-root", default=str(ROOT), help="Generator root (default: this checkout)")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--check", action="store_true",
                      help="Write nothing; exit 1 if a sheet is invalid or show_metrics.sh is out of date")
    mode.add_argument("--add-missing", action="store_true",
                      help="Add rows for metrics and workloads that have none, then rewrite show_metrics.sh")
    args = parser.parse_args(argv)

    root = Path(args.template_root).resolve()
    workbook_path = root / WORKBOOK_NAME
    template = root / TEMPLATE_REL
    try:
        if args.add_missing:
            labels_added, names_added = add_missing_rows(workbook_path)
            print(f"[INFO] Added {labels_added} row(s) to {LABEL_SHEET} and {names_added} row(s) to {NAME_SHEET}"
                  + (f"; rows noted '{AUTO_NOTE}' need a human choice" if labels_added or names_added else ""))
        workbook = load_workbook_file(workbook_path)
        metrics = workbook_metrics(workbook)
        workloads = workbook_workloads(workbook)
        labels = label_rows(workbook)
        names = name_rows(workbook)
        workbook.close()
        errors = validate(metrics, labels) + validate_names(workloads, names)
        if errors:
            for error in errors:
                print(f"[FAIL] {error}", file=sys.stderr)
            return 1
        lines = build_table(labels)
        name_lines = build_name_table(names)
        if not template.is_file():
            raise SheetError(f"template not found: {template}")
        current = template.read_text(encoding="utf-8").replace("\r\n", "\n")
        updated = apply_block(apply_block(current, lines, LABEL_BLOCK), name_lines, NAME_BLOCK)
    except SheetError as exc:
        print(f"[FAIL] {exc}", file=sys.stderr)
        return 1

    overrides = sum(1 for key, _ in lines if ":" in key)
    summary = (f"{len(labels)} metrics in {len({r.workload for r in labels})} workloads -> "
               f"{len(lines) - overrides} key labels + {overrides} workload-specific; "
               f"{len(name_lines)} workload short names")
    if updated == current:
        print(f"[PASS] {TEMPLATE_REL.as_posix()} is up to date ({summary})")
        return 0
    if args.check:
        print(f"[FAIL] {TEMPLATE_REL.as_posix()} is out of date with {LABEL_SHEET}/{NAME_SHEET}; "
              "run python3 scripts/generate_show_metrics.py", file=sys.stderr)
        return 1
    template.write_text(updated, encoding="utf-8", newline="\n")
    print(f"[PASS] Wrote {TEMPLATE_REL.as_posix()} ({summary})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
